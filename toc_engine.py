"""
toc_engine.py — TOC klasifikācija un aprēķini ar TOC-FIN zināšanu bāzes atbalstu
"""
from __future__ import annotations

import json
import re
from dataclasses import dataclass
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from toc_ingest import Leaf, leaf_amounts

FONT_NAME = "Calibri"
FONT_SIZE = 10
KPI_AVG_MODE = "mean_of_months"

GROUPS: dict[str, tuple[str, str]] = {
    "S": ("S", "S"),
    "TVC_MATERIALS": ("Izejvielas un materiāli", "TVC"),
    "TVC_DELIVERY": ("Preču piegādes izdevumi", "TVC"),
    "OE_ADMIN_WAGES": ("Administrācijas algas (+ sociālās iemaksas)", "OE"),
    "OE_WORKER_WAGES": ("Strādnieku algas (+ sociālās iemaksas)", "OE"),
    "OE_MARKETING": ("Mārketings un reklāma", "OE"),
    "OE_PLATFORMS": ("Pārdošanas platformas un citas pārdošanas izm.", "OE"),
    "OE_PREMISES": ("Ražošanas telpu noma un uzturēšana", "OE"),
    "OE_OFFICE_IT": ("Biroja, IT un sakaru izmaksas", "OE"),
    "OE_OTHER": ("Citi saimnieciskie izdevumi", "OE"),
    "EXCLUDED": ("Izslēgtie (memo)", "EXCL"),
}
TVC_KEYS = [k for k, v in GROUPS.items() if v[1] == "TVC"]
OE_KEYS = [k for k, v in GROUPS.items() if v[1] == "OE"]

# Zināšanu bāzes ielāde
KNOWLEDGE_TEXT = ""
for fname in ["toc_knowledge.txt", "toc_knowledge.md", "knowledge.txt"]:
    p = Path(fname)
    if p.exists():
        KNOWLEDGE_TEXT = f"\n\nPAPILDU VADLĪNIJAS NO TOC LITERATŪRAS:\n{p.read_text(encoding='utf-8')}\n"
        break

CLASSIFY_SYSTEM = f"""Tu esi TOC-FIN (Throughput Accounting) vadības lēmumu sistēmas eksperts. Tev jāpiešķir katrai PZA lapai (kontam) TOC grupa.
Skaitļus tu neredzi un nerēķini — tikai klasificē pēc nosaukuma un ceļa.
{KNOWLEDGE_TEXT}
TOC-FIN kontu sanitizācijas un klasifikācijas noteikumi:
- S: neto apgrozījums par pārdotajām precēm un pakalpojumiem.
- TVC: TIKAI un VIENĪGI izdevumi, kas mainās 1:1 ar katru pārdoto vienību (izejvielas TVC_MATERIALS, tiešā preču piegāde TVC_DELIVERY).
- OE: visas darbības un kapacitātes uzturēšanas izmaksas. Algas OBLIGĀTI jāsadala: ražošanas (OE_WORKER_WAGES) un administrācijas (OE_ADMIN_WAGES).
- Pārējie saimnieciskie ieņēmumi (granti, valūtas peļņa) -> OE_OTHER ar is_income=true (tie samazina OE).
- EXCLUDED: "virtuālie grāmatvedības trokšņi" — nolietojums/amortizācija, procentu maksājumi, UIN (peļņas nodoklis), krājumu pārvērtēšana, gatavās produkcijas un nepabeigtās ražošanas (WIP) krājumu izmaiņas.
- ŠAUBU GADĪJUMI: Ja par kādu posteni ir patiesas šaubas (piem. muita, iepakojums, platformu komisijas, ārpakalpojumi), izvēlies visticamāko grupu, iestati confidence='low' un uzdod 1 vienkāršu jautājumu vadības valodā (bez grāmatvedības žargona): vai šis izdevums rodas 1:1 ar katru pārdoto vienību (TVC), vai tiek maksāts par perioda darbības uzturēšanu neatkarīgi no apjoma (OE)."""

CLASSIFY_TOOL = {
    "name": "classify_accounts",
    "description": "Assign a TOC group to every account. Every id must appear exactly once.",
    "input_schema": {
        "type": "object",
        "properties": {
            "items": {
                "type": "array",
                "items": {
                    "type": "object",
                    "properties": {
                        "id": {"type": "string"},
                        "group": {"enum": list(GROUPS)},
                        "is_income": {"type": "boolean"},
                        "confidence": {"enum": ["high", "medium", "low"]},
                        "question": {"type": ["string", "null"],
                                     "description": "Short question in Latvian from economic/decision perspective if confidence is low, else null."},
                    },
                    "required": ["id", "group", "is_income", "confidence"],
                },
            }
        },
        "required": ["items"],
    },
}


def leaves_for_llm(leaves: list[Leaf]) -> str:
    rows = [{"id": l.id, "path": l.path, "label": l.label} for l in leaves]
    return json.dumps(rows, ensure_ascii=False)


def validate_classification(items: list[dict], leaves: list[Leaf]) -> list[str]:
    errs = []
    ids = [i["id"] for i in items]
    want = {l.id for l in leaves}
    missing = sorted(want - set(ids))
    extra = sorted(set(ids) - want)
    dup = sorted({i for i in ids if ids.count(i) > 1})
    if missing:
        errs.append(f"Missing ids: {missing}")
    if extra:
        errs.append(f"Unknown ids: {extra}")
    if dup:
        errs.append(f"Duplicate ids: {dup}")
    bad = [i["id"] for i in items if i.get("group") not in GROUPS]
    if bad:
        errs.append(f"Invalid group for: {bad}")
    return errs


def classify_with_claude(client, model: str, leaves: list[Leaf], max_tries: int = 3) -> list[dict]:
    msgs = [{"role": "user", "content": "Klasificē šīs PZA lapas:\n" + leaves_for_llm(leaves)}]
    items: list[dict] = []
    for _ in range(max_tries):
        resp = client.messages.create(
            model=model, max_tokens=8000, system=CLASSIFY_SYSTEM,
            tools=[CLASSIFY_TOOL], tool_choice={"type": "tool", "name": CLASSIFY_TOOL["name"]},
            messages=msgs,
        )
        tu = next(b for b in resp.content if b.type == "tool_use")
        items = tu.input.get("items", [])
        errs = validate_classification(items, leaves)
        if not errs:
            return items
        msgs += [{"role": "assistant", "content": resp.content},
                 {"role": "user", "content": [{"type": "tool_result", "tool_use_id": tu.id, "is_error": True,
                                               "content": "Fix and resend ALL items. " + " ".join(errs)}]}]
    raise ValueError("Klasifikācija nav pilnīga pēc vairākiem mēģinājumiem: " + " ".join(errs))


@dataclass
class TocResult:
    main: pd.DataFrame
    kpi: pd.DataFrame
    row_kinds: dict[str, str]
    periods: list[str]
    amounts: pd.DataFrame
    checks: list[dict]


def compute_toc(leaves: list[Leaf], mapping: dict[str, dict], period_cols: list[dict],
                reported_result: np.ndarray | None = None) -> TocResult:
    income_ids = {i for i, m in mapping.items() if m.get("is_income")}
    income_ids |= {i for i, m in mapping.items() if m["group"] == "S"}
    A, _ = leaf_amounts(leaves, income_ids)
    ids = [l.id for l in leaves]
    groups = np.array([mapping[i]["group"] for i in ids])

    labels = [p["label"] for p in period_cols]
    month_idx = [j for j, p in enumerate(period_cols) if p["kind"] == "month"]

    def gsum(g: str) -> np.ndarray:
        m = groups == g
        return A[m].sum(axis=0) if m.any() else np.zeros(A.shape[1])

    rows: dict[str, np.ndarray] = {}
    kinds: dict[str, str] = {}
    S = gsum("S"); rows["S"] = S; kinds["S"] = "S"
    tvc_rows = {}
    for k in TVC_KEYS:
        if (groups == k).any():
            tvc_rows[k] = -gsum(k)
    TVC = sum(tvc_rows.values()) if tvc_rows else np.zeros_like(S)
    rows["TVC"] = TVC; kinds["TVC"] = "TVC"
    for k, v in tvc_rows.items():
        rows[GROUPS[k][0]] = v; kinds[GROUPS[k][0]] = "TVC_CHILD"
    T = S - TVC; rows["T"] = T; kinds["T"] = "T"
    oe_rows = {k: -gsum(k) for k in OE_KEYS if (groups == k).any()}
    OE = sum(oe_rows.values()) if oe_rows else np.zeros_like(S)
    rows["OE"] = OE; kinds["OE"] = "OE"
    for k, v in oe_rows.items():
        rows[GROUPS[k][0]] = v; kinds[GROUPS[k][0]] = "OE_CHILD"
    NP = T - OE; rows["NP"] = NP; kinds["NP"] = "NP"
    EXCL = -gsum("EXCLUDED")
    if (groups == "EXCLUDED").any():
        rows[GROUPS["EXCLUDED"][0]] = EXCL; kinds[GROUPS["EXCLUDED"][0]] = "EXCL"

    base = pd.DataFrame(rows, index=labels).T
    if len(month_idx) >= 2:
        mcols = [labels[j] for j in month_idx]
        year = _year_from_labels(mcols)
        base[f"{year} YTD"] = base[mcols].sum(axis=1)
        base[f"{year} AVG YTD"] = base[mcols].mean(axis=1)
    amounts = base.copy()

    main = {}
    for c in amounts.columns:
        main[c] = amounts[c]
        s = amounts.loc["S", c]
        main[f"%|{c}"] = amounts[c] / s if s else np.nan
    main = pd.DataFrame(main)

    def kpi_for(col_vals: pd.Series) -> dict[str, float]:
        S_, T_, OE_ = col_vals["S"], col_vals["T"], col_vals["OE"]
        aw = col_vals.get(GROUPS["OE_ADMIN_WAGES"][0], 0.0)
        ww = col_vals.get(GROUPS["OE_WORKER_WAGES"][0], 0.0)
        div = lambda a, b: a / b if b else np.nan
        return {"PR = T/OE": div(T_, OE_), "T% = T/S": div(T_, S_),
                "Eiro Sales / Eiro algas kopā": div(S_, aw + ww),
                "Eiro Sales / Eiro administrācijas algas": div(S_, aw),
                "Eiro Sales / Eiro strādnieku algas": div(S_, ww)}

    kpi = pd.DataFrame({c: kpi_for(amounts[c]) for c in amounts.columns})
    if len(month_idx) >= 2 and KPI_AVG_MODE == "mean_of_months":
        mcols = [labels[j] for j in month_idx]
        kpi[f"{_year_from_labels(mcols)} AVG YTD"] = kpi[mcols].mean(axis=1)

    checks = []
    tot_ok = np.allclose(NP - EXCL, A.sum(axis=0), atol=1.0)
    checks.append({"Pārbaude": "NP (TOC) − izslēgtās izmaksas = Σ visu lapu", "Statuss": "OK" if tot_ok else "KĻŪDA",
                   "Detaļas": "Katra lapa ir tieši vienā grupā."})
    if reported_result is not None:
        diff = np.abs(NP - EXCL - reported_result).max()
        checks.append({"Pārbaude": "NP (TOC) − izslēgtās izmaksas = faila gala rezultāts",
                       "Statuss": "OK" if diff <= 3.0 else "KĻŪDA",
                       "Detaļas": f"Lielākā starpība: {diff:,.2f}"})
    return TocResult(main, kpi, kinds, list(amounts.columns), amounts, checks)


def _year_from_labels(labels: list[str]) -> str:
    m = re.search(r"(20\d\d)", labels[0])
    return m.group(1) if m else ""


FILLS = {
    "S": ("FFFF00", True, "000000"), "TVC": ("ED7D31", True, "FFFFFF"), "TVC_CHILD": ("F8CBAD", False, "000000"),
    "T": ("00B0F0", True, "000000"), "OE": ("F8CECC", True, "000000"), "OE_CHILD": ("FCE4D6", False, "000000"),
    "NP": ("92D050", True, "000000"), "EXCL": ("D9D9D9", False, "595959"),
}
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def build_excel(res: TocResult, path_or_buf, formulas: bool = True) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "TOC Diagnostika"
    fnt = lambda **kw: Font(name=FONT_NAME, size=FONT_SIZE, **kw)
    amounts = res.amounts
    periods = list(amounts.columns)
    row_names = list(amounts.index)

    ws.cell(1, 1, "Rādītājs / Kontugrupa").font = fnt(bold=True)
    col_of: dict[str, int] = {}
    c = 2
    for p in periods:
        col_of[p] = c
        ws.cell(1, c, p).font = fnt(bold=True)
        ws.cell(1, c + 1, "%").font = fnt(bold=True)
        c += 2
    for cc in range(1, c):
        ws.cell(1, cc).alignment = Alignment(horizontal="center", vertical="center")

    r_of = {name: i + 2 for i, name in enumerate(row_names)}
    months = [p for p in periods if re.match(r"^\d{2}\.\d{4}", p)]
    ytd = next((p for p in periods if p.endswith(" YTD") and "AVG" not in p), None)
    avg = next((p for p in periods if "AVG" in p), None)
    L = get_column_letter
    child_tvc = [n for n in row_names if res.row_kinds[n] == "TVC_CHILD"]
    child_oe = [n for n in row_names if res.row_kinds[n] == "OE_CHILD"]

    for name in row_names:
        r = r_of[name]
        kind = res.row_kinds[name]
        ws.cell(r, 1, name)
        for p in periods:
            cc = col_of[p]
            a, pc = ws.cell(r, cc), ws.cell(r, cc + 1)
            col = L(cc)
            is_ytd, is_avg = p == ytd, p == avg
            val = float(round(amounts.loc[name, p], 2))
            if is_ytd:
                f = "=" + "+".join(f"{L(col_of[m])}{r}" for m in months)
            elif is_avg:
                f = f"=AVERAGE({','.join(f'{L(col_of[m])}{r}' for m in months)})"
            elif kind == "TVC":
                f = f"=SUM({col}{r_of[child_tvc[0]]}:{col}{r_of[child_tvc[-1]]})" if child_tvc else 0
            elif kind == "T":
                f = f"={col}{r_of['S']}-{col}{r_of['TVC']}"
            elif kind == "OE":
                f = f"=SUM({col}{r_of[child_oe[0]]}:{col}{r_of[child_oe[-1]]})" if child_oe else 0
            elif kind == "NP":
                f = f"={col}{r_of['T']}-{col}{r_of['OE']}"
            else:
                f = val
            a.value = f if formulas else val
            pv = res.main.loc[name, f"%|{p}"]
            pc.value = (f"=IF({col}${r_of['S']}=0,\"\",{col}{r}/{col}${r_of['S']})" if formulas
                        else (None if pd.isna(pv) else float(pv)))
            a.number_format = '#,##0;-#,##0;"-"'
            pc.number_format = "0.0%"
            for cell in (a, pc):
                cell.alignment = Alignment(horizontal="right")

        fill, bold, color = FILLS.get(kind, FILLS["OE_CHILD"])
        for cc in range(1, c):
            cell = ws.cell(r, cc)
            cell.fill = PatternFill("solid", start_color=fill, end_color=fill)
            cell.font = fnt(bold=bold, color=color)
            cell.border = BORDER

    k0 = len(row_names) + 4
    ws.cell(k0, 1, "Rādītājs").font = fnt(bold=True)
    kpi_cols = list(res.kpi.columns)
    for j, p in enumerate(kpi_cols):
        ws.cell(k0, 2 + j, p).font = fnt(bold=True)
        ws.cell(k0, 2 + j).alignment = Alignment(horizontal="center")
    S_, T_, OE_ = r_of["S"], r_of["T"], r_of["OE"]
    aw_n = GROUPS["OE_ADMIN_WAGES"][0]
    ww_n = GROUPS["OE_WORKER_WAGES"][0]
    aw = r_of.get(aw_n)
    ww = r_of.get(ww_n)
    kpi_names = list(res.kpi.index)
    for i, kn in enumerate(kpi_names, start=1):
        r = k0 + i
        ws.cell(r, 1, kn)
        for j, p in enumerate(kpi_cols):
            cell = ws.cell(r, 2 + j)
            kv = res.kpi.loc[kn, p]
            avg_of_months = (p == avg and KPI_AVG_MODE == "mean_of_months")
            if avg_of_months:
                f = f"=AVERAGE({L(2)}{r}:{L(1 + len(months))}{r})"
            else:
                a = L(col_of[p])
                den_w = "+".join(f"{a}{x}" for x in (aw, ww) if x)
                f = {
                    "PR = T/OE": f"=IF({a}{OE_}=0,\"\",{a}{T_}/{a}{OE_})",
                    "T% = T/S": f"=IF({a}{S_}=0,\"\",{a}{T_}/{a}{S_})",
                    "Eiro Sales / Eiro algas kopā": f"=IF(({den_w or '0'})=0,\"\",{a}{S_}/({den_w or '1'}))" if den_w else "",
                    "Eiro Sales / Eiro administrācijas algas": f"=IF({a}{aw}=0,\"\",{a}{S_}/{a}{aw})" if aw else "",
                    "Eiro Sales / Eiro strādnieku algas": f"=IF({a}{ww}=0,\"\",{a}{S_}/{a}{ww})" if ww else "",
                }[kn]
            cell.value = (f or None) if formulas else (None if pd.isna(kv) else float(kv))
            cell.number_format = "0.0%" if kn.startswith("T%") else "0.00"
            cell.alignment = Alignment(horizontal="right")
        for cc in range(1, 2 + len(kpi_cols)):
            cell = ws.cell(r, cc)
            cell.fill = PatternFill("solid", start_color="F2F4F4", end_color="F2F4F4")
            cell.font = fnt()
            cell.border = BORDER

    ws.column_dimensions["A"].width = 48
    for cc in range(2, c + 1):
        ws.column_dimensions[L(cc)].width = 12
    ws.freeze_panes = "B2"
    ws.sheet_view.showGridLines = False
    wb.calculation.fullCalcOnLoad = True
    wb.save(path_or_buf)

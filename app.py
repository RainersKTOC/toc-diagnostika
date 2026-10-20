import io
import re
import anthropic
import pandas as pd
import streamlit as st
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

st.set_page_config(page_title="TOC Finanšu Diagnostika", layout="wide", page_icon="📊")

ANTHROPIC_API_KEY = st.secrets.get("ANTHROPIC_API_KEY", "")

# Datu saglabāšana sesijā
if "df_res" not in st.session_state:
    st.session_state.df_res = None
if "output_text" not in st.session_state:
    st.session_state.output_text = None
if "chat_history" not in st.session_state:
    st.session_state.chat_history = []
if "file_text" not in st.session_state:
    st.session_state.file_text = ""

st.title("📊 Finanšu diagnostika pēc TOC principiem")
st.caption("Augšupielādējiet Peļņas un Zaudējumu pārskatu (CSV vai Excel), lai saņemtu noformētu Excel atskaiti un vadības kopsavilkumu.")

if not ANTHROPIC_API_KEY:
    st.error("Nav atrasta ANTHROPIC_API_KEY. Pievienojiet to Streamlit noslēpumos.")
    st.stop()

uploaded_file = st.file_uploader("Izvēlieties PZA failu", type=["xlsx", "xls", "csv"])

def highlight_toc_rows(row):
    label = str(row.iloc[0]).strip().upper()
    if "S (SALES)" in label or label.startswith("S "):
        return ['background-color: #FFFF00; font-weight: bold; color: black;'] * len(row)
    elif "TVC" in label:
        return ['background-color: #F4B084; font-weight: bold; color: black;'] * len(row)
    elif "T (THROUGHPUT)" in label or label == "T":
        return ['background-color: #00B0F0; font-weight: bold; color: white;'] * len(row)
    elif "OE (" in label or label == "OE":
        return ['background-color: #FCE4D6; font-weight: bold; color: black;'] * len(row)
    elif "NP (" in label or "NET PROFIT" in label:
        return ['background-color: #92D050; font-weight: bold; color: black;'] * len(row)
    return ['background-color: #FCE4D6; color: black;'] * len(row)

TOC_SYSTEM_PROMPT = """
Tu esi TOC (Throughput Accounting) finanšu diagnostikas dzinējs.
PZA dati satur gan kopējās kopsavilkuma rindas (piem., 'Pārdotās produkcijas ražošanas pašizmaksa', 'Pārdošanas izmaksas', 'Administrācijas izmaksas', 'Kopā:'), gan arī to detalizētos 'PZA posteņu atšifrējumi'.
UZMANĪBU: NESKAITI DUBULTĀ! Ņem datus TIKAI no atšifrējumu posteņiem!

TOC APRĒĶINU NOTEIKUMI:
1. S (Sales): Neto apgrozījums (100%).
2. TVC (Truly Variable Costs) - TIKAI šie 2 posteņi:
   - Izejvielu un materiālu iepirkšanas un piegādes izd.
   - Preču piegādes izdevumi
3. T (Throughput) = S - TVC.
4. OE (Operating Expenses): Ņem TIKAI atšifrētās darbības izmaksas bez dubultās skaitīšanas:
   - Algas: Strādnieki (Personāla izmaksas-atlīdzība par darbu ražošanā + Sociālās apdrošināšanas obligātās iemaksas ražošanā)
   - Algas: Administrācija (Personāla izmaksas admin + VSAOI admin)
   - Reklāma un mārketings (Reklāmas izmaksas + Mārketinga aktivitātes)
   - Pārdošanas platformas un pārējās pārdošanas izmaksas (Platformu izdevumi + Mājas lapas uzturēšana + Pārējās pārdošanas izmaksas + US TAX)
   - Ražošanas telpas un uzturēšana (Ražošanas telpu noma un uzturēšanas izmaksas + Pārējās ražošanas izmaksas)
   - IT, sakari un birojs (Biroja izmaksas + Telpu noma admin + IT, sakari, bankas)
   - Citi saimnieciskie izdevumi (Apdrošināšana, Ilgspēja, Pārējās saimn. darbības izmaksas mīnus citi saimn. ieņēmumi)
   * PĀRBAUDE: Janvārī OE KOPĀ ir jābūt aptuveni 301 342 EUR, un NP (T - OE) ir jābūt aptuveni +190 110 EUR (kā sākotnējā PZA pirms UIN)!
5. NP (Net Profit) = T - OE.
6. Rādītāji zem NP:
   - PR = T/OE
   - T% = T/S
   - € Sales / € Algas kopā (admin+strādnieki)
   - € Sales / € Administrācijas algas
   - € Sales / € Strādnieku algas

KOLONNU VIRSRAKSTU FORMĀTS:
Augšējā tabulas rindā mēnešus un to procentus norādi tieši šādi (ar gadu un unikāliem nosaukumiem):
Rādītājs / Kontu grupa\t01.2026.\t01.2026. %\t02.2026.\t02.2026. %\t03.2026.\t03.2026. %\t04.2026.\t04.2026. %\t05.2026.\t05.2026. %\t06.2026.\t06.2026. %\t07.2026.\t07.2026. %\t08.2026.\t08.2026. %\t2026 YTD\t2026 YTD %\t2026 AVG YTD\t2026 AVG YTD %

ATBILDES STRUKTŪRA:
Sāc TIEŠI ar ```tsv koda bloku bez ievadvārdiem. Pēc tabulas pievieno 3 īsas sadaļas:
### Kas ir labs un stabils
### Bīstamie signāli
### Vadības fokuss un uzlabojumi
"""

if uploaded_file is not None:
    if st.button("Veikt diagnostiku", type="primary"):
        with st.spinner("Apstrādā datus un veic TOC analīzi..."):
            try:
                if uploaded_file.name.endswith(".csv"):
                    df_raw = pd.read_csv(uploaded_file)
                    st.session_state.file_text = df_raw.to_csv(index=False)
                else:
                    df_raw = pd.read_excel(uploaded_file)
                    st.session_state.file_text = df_raw.to_csv(index=False)

                clean_key = str(ANTHROPIC_API_KEY).strip().strip('"').strip("'")
                client = anthropic.Anthropic(api_key=clean_key)

                response = client.messages.create(
                    model="claude-sonnet-4-6",
                    max_tokens=8000,
                    system=TOC_SYSTEM_PROMPT,
                    messages=[
                        {"role": "user", "content": f"Lūk, PZA dati analīzei:\n\n{st.session_state.file_text}"}
                    ]
                )

                output_text = ""
                for block in response.content:
                    if getattr(block, "type", "") == "text":
                        output_text += block.text
                    elif hasattr(block, "text"):
                        output_text += block.text

                st.session_state.output_text = output_text
                tsv_match = re.search(r"```tsv\n(.*?)```", output_text, re.DOTALL)

                if tsv_match:
                    tsv_data = tsv_match.group(1).strip()
                    st.session_state.df_res = pd.read_csv(io.StringIO(tsv_data), sep="\t")
                    st.session_state.chat_history = []

            except Exception as e:
                st.error(f"Kļūda apstrādē: {str(e)}")

if st.session_state.df_res is not None:
    st.success("Diagnostika ir aktīva!")

    excel_buffer = io.BytesIO()
    with pd.ExcelWriter(excel_buffer, engine="openpyxl") as writer:
        st.session_state.df_res.to_excel(writer, index=False, sheet_name="TOC Diagnostika")
        ws = writer.sheets["TOC Diagnostika"]

        fill_yellow = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
        fill_orange = PatternFill(start_color="F4B084", end_color="F4B084", fill_type="solid")
        fill_cyan = PatternFill(start_color="00B0F0", end_color="00B0F0", fill_type="solid")
        fill_peach = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
        fill_green = PatternFill(start_color="92D050", end_color="92D050", fill_type="solid")

        thin_border = Border(
            left=Side(style='thin', color='BFBFBF'),
            right=Side(style='thin', color='BFBFBF'),
            top=Side(style='thin', color='BFBFBF'),
            bottom=Side(style='thin', color='BFBFBF')
        )

        for col_num in range(1, len(st.session_state.df_res.columns) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = Font(name="Calibri", size=10, bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for r_idx, row in enumerate(st.session_state.df_res.itertuples(index=False), start=2):
            label = str(row[0]).strip().upper()
            row_fill = fill_peach
            row_font = Font(name="Calibri", size=10)

            if "S (SALES)" in label or label.startswith("S "):
                row_fill = fill_yellow
                row_font = Font(name="Calibri", size=10, bold=True)
            elif "TVC" in label:
                row_fill = fill_orange
                row_font = Font(name="Calibri", size=10, bold=True)
            elif "T (THROUGHPUT)" in label or label == "T":
                row_fill = fill_cyan
                row_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
            elif "NP (" in label or "NET PROFIT" in label:
                row_fill = fill_green
                row_font = Font(name="Calibri", size=10, bold=True)
            elif "OE (" in label:
                row_font = Font(name="Calibri", size=10, bold=True)

            for c_idx in range(1, len(st.session_state.df_res.columns) + 1):
                c = ws.cell(row=r_idx, column=c_idx)
                c.fill = row_fill
                c.font = row_font
                c.border = thin_border

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

    st.download_button(
        label="📥 Lejupielādēt noformēto Excel failu (.xlsx)",
        data=excel_buffer.getvalue(),
        file_name="TOC_Finansu_Diagnostika.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    st.subheader("Finanšu pārskata tabula")
    styled_df = st.session_state.df_res.style.apply(highlight_toc_rows, axis=1)
    st.dataframe(styled_df, use_container_width=True)

    st.subheader("Vadības analīze un secinājumi")
    clean_text = re.sub(r"```tsv.*?```", "", st.session_state.output_text, flags=re.DOTALL).strip()
    st.markdown(clean_text)

    st.markdown("---")
    st.subheader("💬 Uzdod jautājumu par šo TOC analīzi")

    for msg in st.session_state.chat_history:
        with st.chat_message(msg["role"]):
            st.markdown(msg["content"])

    user_query = st.chat_input("Ieraksti jautājumu par pārskata datiem...")
    if user_query:
        st.session_state.chat_history.append({"role": "user", "content": user_query})
        with st.chat_message("user"):
            st.markdown(user_query)

        with st.chat_message("assistant"):
            with st.spinner("Domā..."):
                clean_key = str(ANTHROPIC_API_KEY).strip().strip('"').strip("'")
                client = anthropic.Anthropic(api_key=clean_key)
                
                chat_context = f"""
                Konteksts ar TOC analīzi un datiem:
                {st.session_state.output_text}
                
                Lietotāja jautājums:
                {user_query}
                
                Atbildi profesionāli, kodolīgi un balstoties uz šiem datiem un TOC principiem.
                """
                
                chat_resp = client.messages.create(
                    model="claude-sonnet-4-6",
                    max_tokens=1500,
                    messages=[{"role": "user", "content": chat_context}]
                )
                
                reply = ""
                for block in chat_resp.content:
                    if getattr(block, "type", "") == "text":
                        reply += block.text
                    elif hasattr(block, "text"):
                        reply += block.text
                        
                st.markdown(reply)
                st.session_state.chat_history.append({"role": "assistant", "content": reply})
